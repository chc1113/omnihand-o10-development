from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "O10_灵巧手二次开发阶段工作文档.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size=10.5, bold=False, color="000000") -> None:
    run.font.name = "Noto Sans CJK SC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=5, line=1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line


def add_text(doc, text, style=None, bold_prefix=None) -> None:
    p = doc.add_paragraph(style=style)
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)


def add_bullets(doc, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, space_after=2)
        r = p.add_run(item)
        set_run_font(r, size=10.2)


def add_table(doc, headers, rows, widths=None, font_size=9.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(hdr[i], "44546A")
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line=1.0)
        r = p.add_run(str(h))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F2F4F7")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, space_after=0, line=1.05)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, text) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=9.2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    for name in ("Normal", "Body Text"):
        st = styles[name]
        st.font.name = "Noto Sans CJK SC"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        st.font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.2)):
        st = styles[name]
        st.font.name = "Noto Sans CJK SC"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        ppr = st._element.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is not None:
            ppr.remove(pbdr)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(p, space_after=6, line=1.05)
    r = p.add_run(text)
    set_run_font(r, size={1: 15, 2: 12.5, 3: 11.2}[level], bold=True)
    return p


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("O10 灵巧手二次开发  |  阶段工作文档")
    set_run_font(fr, size=8.5, color="666666")

    title = doc.add_paragraph(style="Title")
    title_ppr = title._p.get_or_add_pPr()
    title_pbdr = title_ppr.find(qn("w:pBdr"))
    if title_pbdr is not None:
        title_ppr.remove(title_pbdr)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, space_after=8, line=1.0)
    r = title.add_run("O10 灵巧手二次开发阶段工作文档")
    set_run_font(r, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(sub, space_after=16, line=1.0)
    r = sub.add_run("右手 USB 视觉追踪与触觉交互项目准备")
    set_run_font(r, size=12, color="44546A")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, space_after=18, line=1.0)
    r = meta.add_run(f"测试环境：Linux x86_64  |  SDK 1.1.5  |  文档日期：{date.today().isoformat()}")
    set_run_font(r, size=9.5, color="666666")

    heading(doc, "一 文档目的与当前结论", 1)
    add_text(doc, "本文档汇总当前 O10 灵巧手、USB 通讯链路和本机摄像头的实测结果，并将后续工作拆分为单视觉项目、单触觉项目以及视觉触觉混合项目三条路线。当前开发范围收敛为右手单手项目，左手仅作为已完成的对照测试，不进入后续控制程序。")
    add_text(doc, "结论先行：右手 USB 的位置与状态反馈稳定可用；右手降采样触觉在拇指指尖施压时已实测出现非零阵列，释放后回零，因此可以支持接触检测和交互事件。USB 不支持原始高分辨率触觉帧，暂时不具备高频精确力控条件。摄像头可用，适合先做固定机位下的二维目标追踪。")
    add_bullets(doc, [
        "当前首选平台：右手 O10 + USB CDC + 本机 /dev/video0 摄像头。",
        "推荐第一版频率：关节位置 50 Hz，电流/温度/错误 5～10 Hz，降采样触觉 10 Hz，视觉采集 30 Hz。",
        "尚未进行电机动作测试；后续首次动作必须设置低速、小幅度、限位、电流保护和立即断电条件。",
        "当前没有 CANFD 硬件，原始触觉和高带宽力控路线暂不纳入第一版实现。",
    ])

    heading(doc, "二 项目与硬件配置", 1)
    add_table(doc, ["项目项", "当前配置", "状态", "说明"], [
        ("独立 SDK", "o10_usb_phase1/sdk/agillink_omnihand_sdk-main", "已完成", "从官方压缩包重新解压，未修改原有演示 SDK"),
        ("Python 环境", ".venv，Python 3.14.6", "已完成", "可导入 omnihand，SDK 版本 1.1.5"),
        ("右手 USB", "/dev/ttyACM0；by-id 自动发现", "已验证", "当前项目主设备，设备 ID 1"),
        ("摄像头", "/dev/video0，HD Webcam", "已验证", "可采集；/dev/video1 不是采集节点"),
        ("CANFD", "无 can0/can1，无 USB-CANFD 适配器", "缺少硬件", "原始触觉链路暂时无法验证"),
    ], widths=[1.25, 2.2, 0.9, 2.6])
    add_text(doc, "安全边界：本阶段所有 USB 长测均为只读调用，没有向 joint_cmd、joint_mix_control_cmd、位置写入、电流阈值、功率或控制模式接口发送命令。串口测试遵循单进程占用原则。")

    heading(doc, "三 硬件测试数据", 1)
    heading(doc, "3.1 左手对照测试", 2)
    add_table(doc, ["测试项", "时长与目标", "实际结果", "延迟与异常"], [
        ("位置读取", "5 分钟，50 Hz", "15000/15000；49.99999 Hz", "调用均值 0.444 ms；P99 0.742 ms；最大 0.983 ms；0 失败"),
        ("全量状态（不含触觉）", "5 分钟，10 Hz", "3000/3000；10.00000 Hz", "周期均值 4.043 ms；P99 4.726 ms；最大 6.469 ms；各项 0 失败"),
        ("降采样触觉", "20 秒，5 Hz", "手指 16 点，掌心/手背 25 点", "手指与掌心全零；手背仅出现单点 1 的噪声"),
        ("原始触觉", "单次探针", "thumb data=[]；all=[]", "SDK/ROS2 明确提示 USB 不支持原始触觉"),
    ], widths=[1.45, 1.35, 2.0, 2.15])
    add_text(doc, "左手结论：位置、角度、电流、温度和错误反馈可用，但当前左手没有得到有效压力触觉证据，不作为后续触觉项目执行设备。")

    heading(doc, "3.2 右手 USB 基线与压力测试", 2)
    add_table(doc, ["测试项", "时长与目标", "实际结果", "结论"], [
        ("设备信息", "USB 探针", "O10；ID 1；HW 3.0.0；SW 1.2.7；10 DOF", "设备身份与 USB 初始化正常"),
        ("位置读取", "5 分钟，50 Hz", "15000/15000；49.99999 Hz", "P99 0.772 ms；最大 0.955 ms；0 失败"),
        ("全量状态与降采样触觉", "5 分钟，10 Hz", "3000/3000；10.00000 Hz", "周期均值 11.064 ms；P99 11.787 ms；最大 13.089 ms；0 失败"),
        ("空载触觉基线", "同一全量测试", "7 区域均返回正确长度，3000 组均全零", "接口可访问，空载无响应"),
        ("拇指指尖施压", "30 秒，10 Hz", "300/300 周期拇指非零；最高值约 150", "拇指 16 点阵列有稳定压力响应"),
        ("拇指释放", "10 秒，10 Hz", "100/100 周期拇指全零", "释放后回零，具备接触/释放事件基础"),
    ], widths=[1.55, 1.35, 2.45, 2.0])
    add_text(doc, "拇指施压代表性数组： [2, 2, 80, 92, 145, 146, 132, 132, 66, 23, 0, 0, 0, 0, 0, 0]。同一测试中食指、中指、无名指、小拇指、掌心和手背保持全零。当前数据是 SDK 降采样值，不能直接换算为牛顿力，正式项目需要零点、阈值、滞回和标定。")

    heading(doc, "3.3 摄像头与通信能力", 2)
    add_table(doc, ["硬件/接口", "检测结果", "对项目的意义"], [
        ("/dev/video0", "HD Webcam；640x480@30（JPEG/YUY2）；最高 1280x720@10", "可用于第一版固定机位二维追踪"),
        ("JPEG 采集", "640x480，300 帧约 10.49 秒，约 28.6 fps", "满足 30 Hz 视觉采集目标"),
        ("YUY2 采集", "640x480，150 帧约 5.53 秒，约 27.1 fps", "可作为低延迟原始帧路径"),
        ("/dev/video1", "同一 UVC 设备的非采集节点，GStreamer 报 not a capture device", "不纳入软件输入"),
        ("USB 帧", "EE AA | device_id(2B LE) | length(1B) | command+payload | crc(2B)", "可在桥接层保留协议日志"),
        ("CANFD", "未发现 can0/can1 或适配器", "不能验证原始触觉多帧和高频力控"),
    ], widths=[1.45, 3.15, 2.75])

    heading(doc, "四 统一软件数据链路", 1)
    add_text(doc, "建议建立右手专用的 o10_usb_bridge，将 USB 访问、时间戳、状态发布和记录与视觉算法解耦。视觉程序不直接创建 SDK Hand 对象，避免串口抢占和读取频率互相影响。")
    add_code(doc, "摄像头 30 Hz -> 目标检测/跟踪 -> 目标状态\n右手 USB 50 Hz -> 关节状态\n右手 USB 5~10 Hz -> 电流/温度/错误\n右手 USB 10 Hz -> 七区域降采样触觉\n以上数据统一时间戳 -> 记录/决策/安全控制")
    add_text(doc, "建议的数据模型包含 timestamp、side、joint_position[10]、joint_angle[10]、joint_current[10]、joint_temperature[10]、joint_error[10] 和 tactile[thumb/index/middle/ring/little/palm/dorsum]。触觉字段同时保存 raw_uint8、区域最大值、区域总和、非零点数量和接触状态；不要把没有数据的区域伪造为零，应增加 tactile_available 和 data_quality 字段。")
    add_text(doc, "记录格式优先采用 JSONL 或 ROS 2 bag，并保存摄像头帧时间戳、视觉目标状态和右手状态。这样可以先离线回放视觉与触觉数据，再开放真实控制。")

    heading(doc, "五 单视觉项目方案", 1)
    add_text(doc, "目标：仅使用本机摄像头完成目标检测、跟踪和跟随意图输出，第一阶段不驱动右手。推荐先使用颜色目标、ArUco 或 AprilTag，建立稳定的二维基线，再替换为人手或通用目标模型。")
    add_table(doc, ["阶段", "工作内容", "输出与验收"], [
        ("V0 采集基线", "固定 /dev/video0；选择 640x480@30；记录帧时间戳和丢帧", "连续 10 分钟采集，无明显断流；实测帧率接近 30 Hz"),
        ("V1 目标检测", "输出目标框、中心、面积、置信度；加入 ROI 和异常过滤", "目标进入/离开画面时状态正确，误检可追踪"),
        ("V2 跟踪滤波", "卡尔曼或指数滤波；定义丢失超时和重新捕获", "目标中心抖动受控，丢失后不产生虚假动作"),
        ("V3 跟随意图", "将图像偏差映射为左/右、靠近/远离、停止建议", "只显示建议动作，不发 joint_cmd"),
        ("V4 低风险控制", "单关节、小幅度、低速、限位和急停", "人工在场下完成可重复跟随，错误时立即停止"),
    ], widths=[1.15, 3.3, 2.9])
    add_text(doc, "边界：单目摄像头天然缺少可靠深度。第一版应把跟随定义为二维图像空间跟随；若任务要求真实距离、空间抓取或遮挡下稳定跟踪，需要增加深度摄像头、双目或外部定位。")

    heading(doc, "六 单触觉项目方案", 1)
    add_text(doc, "目标：使用右手七个触觉区域实现接触、持续压力、压力增强、释放和过载保护等交互，不依赖摄像头。当前 USB 可用的是约 10 Hz 的降采样数据，适合交互状态机，不适合高频精确力控。")
    add_table(doc, ["阶段", "工作内容", "输出与验收"], [
        ("T0 区域盘点", "逐个轻压拇指、五指、掌心、手背，记录每区域 16/25 点响应", "形成区域映射和基线记录；确认各区域非零响应"),
        ("T1 零点与阈值", "空载采样；计算均值、峰值、噪声；设置接触/释放双阈值", "压力变化不抖动；释放后能回到空载状态"),
        ("T2 特征提取", "区域最大值、总和、非零点数、变化速度、持续时间", "输出统一 TactileState"),
        ("T3 交互状态机", "无接触、接触开始、持续、增强、过载、释放", "事件可记录、可回放、可重复触发"),
        ("T4 与控制联动", "接触暂停跟随、减速、停止或确认抓取", "触觉事件优先级高于视觉跟随"),
    ], widths=[1.15, 3.3, 2.9])
    add_text(doc, "建议的触觉状态机：NO_CONTACT -> CONTACT_START -> CONTACT_HOLD -> PRESSURE_HIGH -> OVERLOAD；压力下降后进入 RELEASE，再回到 NO_CONTACT。每次状态切换都记录时间戳、区域和触觉特征。")
    add_text(doc, "标定注意：SDK 数值单位目前只能按原始/降采样等级使用，不能直接标称克或牛顿。若要得到力值，需要已知载荷、多点标定和温漂复测；若要使用原始高分辨率阵列，需要补充 CANFD 硬件。")

    heading(doc, "七 视觉触觉混合项目方案", 1)
    add_text(doc, "目标：摄像头负责目标位置和运动趋势，右手触觉负责接触确认、安全减速和释放判断，构成视觉跟随与触觉交互的单手闭环。推荐先做监控闭环，再逐步开放动作闭环。")
    add_table(doc, ["层级", "职责", "建议实现"], [
        ("感知层", "摄像头与七区域触觉采集", "独立线程、统一单调时钟和墙上时间戳"),
        ("融合层", "对齐视觉目标与触觉事件", "允许视觉帧和触觉帧异步到达，使用最近时间戳匹配"),
        ("决策层", "跟随、接触、暂停、释放、异常", "有限状态机，触觉安全事件优先"),
        ("控制层", "生成右手动作建议或低风险命令", "先建议后执行；执行时限制幅度、速度和频率"),
        ("记录层", "复盘和故障定位", "视频、状态、触觉、命令和错误统一记录"),
    ], widths=[1.05, 2.4, 3.9])
    add_bullets(doc, [
        "视觉锁定目标且无触觉：允许继续跟随。",
        "视觉锁定目标且检测到接触：减速或暂停，等待决策。",
        "触觉压力持续升高或超过阈值：立即停止增力，必要时回退。",
        "视觉目标丢失但触觉仍持续：保持安全姿态，不继续追踪运动。",
        "触觉释放且目标重新可见：允许恢复跟随，但需要滞回和重新确认。",
    ])
    add_text(doc, "混合项目的关键验收不是单个传感器有数据，而是时间对齐、状态切换和安全优先级可重复。第一版可先用录制数据离线验证状态机，再做在线控制。")

    heading(doc, "八 推荐实施顺序", 1)
    add_table(doc, ["顺序", "里程碑", "完成标志"], [
        ("1", "右手 USB 桥接与记录", "单进程稳定读取；状态和触觉统一时间戳；可回放"),
        ("2", "单视觉监控版", "实时画面、目标框、轨迹、丢失报警，不驱动手"),
        ("3", "单触觉标定版", "七区域逐区验证；阈值、滞回和触觉事件稳定"),
        ("4", "视觉触觉离线融合", "录制数据上验证接触、暂停、释放和异常逻辑"),
        ("5", "右手低风险动作", "单关节、小幅度、低速，人工急停可用"),
        ("6", "视觉触觉在线闭环", "目标跟随、接触减速/停止、释放恢复可重复"),
    ], widths=[0.65, 2.45, 4.25])

    heading(doc, "九 风险与配置要求", 1)
    add_table(doc, ["风险/限制", "影响", "应对"], [
        ("USB 原始触觉不支持", "不能做高频精确力控", "第一版只做降采样接触交互；后续补 CANFD"),
        ("触觉数值未标定", "不能直接输出物理力值", "做零点、阈值、滞回和已知载荷标定"),
        ("单目无可靠深度", "三维跟随和距离控制受限", "先做二维；需要时增加深度/双目"),
        ("尚未做电机动作测试", "控制参数和安全边界未知", "人工在场，低速小幅度，急停和限位先行"),
        ("ROS2 节点配置约束", "by-id 路径中的连字符会导致节点名非法", "节点 smoke test 使用 /dev/ttyACM0，Python 桥接使用 by-id"),
        ("串口独占", "并行程序会互相抢占请求/响应", "统一由桥接层管理，禁止多个访问进程"),
    ], widths=[1.75, 2.45, 3.15])
    add_text(doc, "运行前置条件：确认右手供电稳定、手指周围空旷、USB by-id 路径有效、没有 Toolbox 或其他程序占用串口；所有动作测试均应保留立即断电手段。")

    heading(doc, "十 现有记录与项目文件", 1)
    add_bullets(doc, [
        "右手位置长测：records/right_positions_50hz_5min.jsonl",
        "右手全量长测：records/right_full_10hz_5min.jsonl",
        "右手拇指施压：records/right_thumb_pressure_10hz_30s.jsonl",
        "右手拇指释放：records/right_thumb_release_10hz_10s.jsonl",
        "右手评估：reports/right_usb_assessment.md",
        "当前 ROS2 配置：config/omnihand_2025_usb_readonly.yaml",
        "只读监测程序：scripts/usb_readonly_monitor.py",
    ])
    add_text(doc, "下一步建议直接进入里程碑 1：实现右手 USB 桥接节点和统一记录格式，然后分别完成单视觉监控版与全区域触觉标定版。完成这两个独立基线后，再做视觉触觉融合，能够清楚定位问题来自摄像头、触觉还是控制逻辑。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
