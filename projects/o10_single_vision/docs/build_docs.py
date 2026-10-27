from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', size='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); mar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def setup(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Noto Sans CJK SC'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.15
    for name,size in [('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
        st=styles[name]; st.font.name='Noto Sans CJK SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_before=Pt(10 if name!='Title' else 0); st.paragraph_format.space_after=Pt(5)
        if name == 'Title':
            ppr = st._element.get_or_add_pPr()
            border = ppr.find(qn('w:pBdr'))
            if border is not None:
                ppr.remove(border)

def add_table(doc, rows):
    table=doc.add_table(rows=1, cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for j,val in enumerate(rows[0]):
        c=table.rows[0].cells[j]; c.text=val; set_cell_shading(c,'5B6573'); set_cell_border(c); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9.5)
    for i,row in enumerate(rows[1:]):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            cells[j].text=val; set_cell_border(cells[j]); set_cell_margins(cells[j]); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i%2==1: set_cell_shading(cells[j],'F2F4F6')
            for p in cells[j].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    for row in table.rows:
        trPr=row._tr.get_or_add_trPr(); rep=trPr.find(qn('w:tblHeader'))
        if rep is None: rep=OxmlElement('w:tblHeader'); trPr.append(rep)
    doc.add_paragraph()

def build(md_path, out):
    doc=Document(); setup(doc)
    lines=md_path.read_text(encoding='utf-8').splitlines(); i=0; first=True
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title'); p.add_run(line[2:].strip()); first=False; i+=1; continue
        if line.startswith('## '): doc.add_paragraph(line[3:].strip(),style='Heading 1'); i+=1; continue
        if line.startswith('### '): doc.add_paragraph(line[4:].strip(),style='Heading 2'); i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if all(set(x)<=set('-: ') for x in vals): i+=1; continue
                rows.append(vals); i+=1
            if rows: add_table(doc,rows)
            continue
        if line.startswith('```'):
            i+=1; buf=[]
            while i<len(lines) and not lines[i].startswith('```'): buf.append(lines[i]); i+=1
            i+=1; p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); r=p.add_run('\n'.join(buf)); r.font.name='DejaVu Sans Mono'; r.font.size=Pt(9); continue
        if line.startswith('- ') or line[:2].isdigit() and line[2:3]=='.':
            text=line[2:] if line.startswith('- ') else line[3:]
            p=doc.add_paragraph(style='List Bullet' if line.startswith('- ') else 'List Number'); p.add_run(text); i+=1; continue
        p=doc.add_paragraph(); p.add_run(line); i+=1
    doc.save(out)

if __name__=='__main__':
    py=Path('/home/chc/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3')
    build(ROOT/'工作总结_单视觉项目.md', ROOT/'工作总结_单视觉项目.docx')
    build(ROOT/'项目说明与操作手册_单视觉.md', ROOT/'项目说明与操作手册_单视觉.docx')
